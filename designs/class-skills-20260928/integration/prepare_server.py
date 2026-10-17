from pathlib import Path
import json, shutil
from openpyxl import load_workbook
ROOT=Path('E:/work/xuanming-server-mmo')
OUT=Path('E:/work/image/designs/class-skills-20260928/integration/server')
CAT=json.loads(Path('E:/work/image/designs/class-skills-20260928/catalog.json').read_text(encoding='utf-8'))
def read(p): return (ROOT/p).read_text(encoding='utf-8-sig')
def write(p,s):
    t=OUT/p;t.parent.mkdir(parents=True,exist_ok=True);t.write_text(s,encoding='utf-8')
def replace(s,a,b):
    if a not in s: raise Exception('missing '+a[:100])
    return s.replace(a,b,1)

s=read('proto/common/component/player_skill_comp.proto')
s=replace(s,'uint32 skill_table_id = 2;', 'uint32 skill_table_id = 2;\n  uint32 rank = 3; // 职业技能等级；0为未学习，历史通用技能读取时兼容为1。')
write('proto/common/component/player_skill_comp.proto',s)
s=read('proto/scene/player_skill.proto')
s=replace(s,'PlayerSkillListComp skill_list = 2;','''PlayerSkillListComp skill_list = 2;
  repeated SkillCatalogEntry catalog = 3;
  uint32 player_level = 4;
  uint32 class_id = 5;
  uint64 currency_balance = 6; // 金币，沿用 CurrencySystem 的 0 号币种。''')
extra='''
// 面板投影；技能数值真源仍是 Skill 配置表，客户端不得回传这些派生值。
message SkillCatalogEntry {
  uint32 skill_table_id = 1;
  string name = 2;
  string description = 3;
  string icon_key = 4;
  uint32 class_id = 5;
  string category = 6;
  bool passive = 7;
  uint32 target_team = 8; // 1敌方、2友方、3自己
  uint32 max_targets = 9;
  uint32 unlock_level = 10;
  uint32 max_rank = 11;
  uint32 rank = 12;
  uint32 rank_limit = 13;
  double mana_percent = 14;
  uint32 cooldown_rounds = 15;
  uint64 next_rank_cost = 16;
  string upgrade_block_reason = 17;
}
message UpgradeSkillRequest {
  uint32 skill_table_id = 1;
  uint32 expected_rank = 2; // 乐观比较：重放旧请求不重复扣款。
  uint32 count = 3; // 仅1或10；批量在等级、上限、余额处停止。
}
message UpgradeSkillResponse {
  TipInfoMessage error_message = 1;
  ListSkillsResponse state = 2;
  uint32 upgraded_count = 3;
  string stop_reason = 4; // 部分成功也返回明确停止原因。
}
'''
s=replace(s,'service SceneSkillClientPlayer {',extra+'\nservice SceneSkillClientPlayer {')
s=replace(s,'rpc ListSkills (ListSkillsRequest) returns (ListSkillsResponse);','rpc ListSkills (ListSkillsRequest) returns (ListSkillsResponse);\n  rpc UpgradeSkill (UpgradeSkillRequest) returns (UpgradeSkillResponse);')
write('proto/scene/player_skill.proto',s)
s=read('proto/battle/battle_data.proto')
s=replace(s,'uint64 caster_id = 5;','uint64 caster_id = 5;\n  uint32 applied_round = 6; // 新职业增益不在施加回合末扣持续时间。')
s=replace(s,'uint32 gender = 19;','uint32 gender = 19;\n  map<uint32, uint32> skill_ranks = 20; // 已学主动和被动，等级由scene校验。')
s=replace(s,'uint32 gender = 26;','uint32 gender = 26;\n  map<uint32, uint32> skill_ranks = 27;')
write('proto/battle/battle_data.proto',s)

fields=[('uint32','class_id'),('string','name'),('string','description'),('string','icon_key'),('string','category'),('uint32','unlock_level'),('uint32','max_rank'),('uint32','target_team'),('uint32','max_targets'),('double','mana_percent'),('uint32','cooldown_rounds'),('double','heal_multiplier'),('double','heal_per_level'),('double','execute_heal_threshold'),('double','execute_heal_bonus'),('double','restore_mana_percent'),('uint32','dispel_count'),('int32','initiative'),('string','passive_stat'),('double','passive_percent'),('uint32','initial_learned')]
s=read('data/schema/skill_table.proto')
pos=s.rfind('}')
s=s[:pos]+'\n  // 四职业技能：0值保持历史通用技能兼容。百分比统一0–1比例。\n'+''.join(f'  {typ} {name} = {i};\n' for i,(typ,name) in enumerate(fields,26))+s[pos:]
write('data/schema/skill_table.proto',s)
def workbook(name):
    book=load_workbook(ROOT/('data/'+name+'.xlsx'));sheet=book.active
    return book,sheet,{str(c.value):c.column for c in sheet[1] if c.value is not None}
def append(sheet,cols,values):
    r=sheet.max_row+1
    for k,v in values.items():
        if isinstance(k,int): sheet.cell(r,k,v); continue
        if k not in cols: raise ValueError(k)
        sheet.cell(r,cols[k],v)
def save(book,name):
    p=OUT/('data/'+name+'.xlsx');p.parent.mkdir(parents=True,exist_ok=True);book.save(p)
book,ws,cols=workbook('Skill')
for typ,key in fields:
    if key not in cols: cols[key]=ws.max_column+1;ws.cell(1,cols[key],key)
buffbook,bws,bcols=workbook('Buff')
for item in CAT['skills']:
    def col(prefix,part=None):
        candidates=[k for k in cols if k.startswith(prefix) and (part is None or part in k)]
        return candidates[0]
    v={k:0 for k in cols}
    v.update({'id':item['id'],'damage':str(item['baseDamage']),'damage_type':1 if item['damageType']=='physical' else 0,'attack_multiplier':item['attackMultiplier'],'class_id':item['classId'],'name':item['name'],'description':item['description'],'icon_key':f"UI/Skills/{item['id']}",'category':item['category'],'unlock_level':item['unlockLevel'],'max_rank':10,'target_team':{'enemy':1,'ally':2,'self':3}[item['targetTeam']],'max_targets':item['maxTargets'],'mana_percent':item['manaPercent'],'cooldown_rounds':item['cooldownRounds'],'heal_multiplier':item['healMultiplier'],'heal_per_level':item['healPerLevel'],'execute_heal_threshold':item['executeHealThreshold'],'execute_heal_bonus':item['executeHealBonus'],'restore_mana_percent':item['restoreManaPercent'],'dispel_count':item['dispelCount'],'initiative':item['initiative'],'passive_stat':item['passiveStat'] if item['type']=='passive' else '', 'passive_percent':item['passivePercent'],'initial_learned':int(item['id']%100 in (1,6))})
    for k in cols:
        if k.startswith('skill_type'): v[k]=None
        if k.startswith('targeting_mode'): v[k]=None
    v[col('skill_type')]=0 if item['type']=='passive' else 1
    v[col('targeting_mode')]=1
    if item['buffKind']!='none':
        v[col('effect')]=10000+item['id']
        bv={'id':10000+item['id'],'buff_type':{'defense':11,'physicalAttack':10,'speed':13}[item['buffKind']],'max_layer':1,'duration':item['durationRounds']*6,'no_caster':1}
        tags=[k for k in bcols if k.startswith('tag')]
        bv[tags[0]]='class_skill'
        if item['buffPercent']<0: bv[bcols['tag']+1]='debuff'
        bv[next(k for k in bcols if k.startswith('interval_effect'))]=item['buffPercent']
        append(bws,bcols,bv)
    append(ws,cols,v)
save(book,'Skill');save(buffbook,'Buff')

rules='''#pragma once
#include <algorithm>
#include <cstdint>
// 学习、面板、战斗共用规则；所有费用用金币，等级只由服务端提交。
namespace classskill {
inline constexpr uint32_t kMaxRank = 10;
inline uint32_t RankLimit(uint32_t level, uint32_t unlock) {
    return level < unlock ? 0u : std::min(kMaxRank, 1u + (level - unlock) / 2u);
}
inline uint64_t UpgradeCost(uint32_t nextRank, uint32_t unlock) {
    return 100ull * nextRank * nextRank + 10ull * unlock;
}
inline double RankFactor(uint32_t rank) { return 1.0 + .02 * (std::clamp(rank,1u,kMaxRank)-1); }
inline double PassiveFactor(uint32_t rank) { return 1.0 + .05 * (std::clamp(rank,1u,kMaxRank)-1); }
}
'''
write('cpp/libs/services/battle/system/class_skill_rules.h',rules)
print('staged proto, schema, tables and rules')
